"""Tests for document intelligence nodes."""

from __future__ import annotations

import io
import json
import subprocess
import sys
from types import SimpleNamespace

import pytest

import noodle_nodes  # noqa: F401 - registers nodes
from noodle.artifacts import LocalArtifactStore, is_artifact_ref
from noodle.context import artifact_store, current_node_id
from noodle.datasets import is_dataset_ref


@pytest.fixture
def store_ctx(tmp_path):
    store = LocalArtifactStore(tmp_path, run_id="doc-intel-test-run")
    a = artifact_store.set(store)
    n = current_node_id.set("doc-intel-test-node")
    yield store
    current_node_id.reset(n)
    artifact_store.reset(a)


def test_pdf_extract_text_raises_without_input(store_ctx) -> None:
    from noodle_nodes.document_intelligence import pdf_extract_text
    with pytest.raises(ValueError, match="input is required"):
        pdf_extract_text(input=None)


def test_pdf_extract_text_raises_missing_package(store_ctx, monkeypatch) -> None:
    """Lazy import raises RuntimeError with install instructions when pdfplumber absent."""
    import builtins
    real_import = builtins.__import__

    def mock_import(name, *args, **kwargs):
        if name == "pdfplumber":
            raise ImportError("No module named 'pdfplumber'")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", mock_import)
    # Write a fake artifact so we pass the None check
    from noodle.artifacts import write_bytes
    from noodle_nodes.document_intelligence import pdf_extract_text
    fake_ref = write_bytes(b"fake", name="test.pdf", content_type="application/pdf")
    with pytest.raises(RuntimeError, match="pdfplumber"):
        pdf_extract_text(input=fake_ref)


def test_pdf_extract_text_raises_for_scanned_pdf(store_ctx) -> None:
    """PDF with no text layer raises ValueError with helpful message."""
    pytest.importorskip("pdfplumber")
    from unittest.mock import MagicMock, patch

    from noodle.artifacts import write_bytes
    from noodle_nodes.document_intelligence import pdf_extract_text

    fake_ref = write_bytes(b"fake", name="test.pdf", content_type="application/pdf")

    mock_page = MagicMock()
    mock_page.extract_text.return_value = ""
    mock_page.page_number = 1

    mock_pdf_ctx = MagicMock()
    mock_pdf_ctx.__enter__ = lambda s: s
    mock_pdf_ctx.__exit__ = MagicMock(return_value=False)
    mock_pdf_ctx.pages = [mock_page]

    with patch("pdfplumber.open", return_value=mock_pdf_ctx):
        with pytest.raises(ValueError, match="scanned"):
            pdf_extract_text(input=fake_ref)


def test_pdf_extract_text_returns_text_and_artifact(store_ctx) -> None:
    pytest.importorskip("pdfplumber")
    from unittest.mock import MagicMock, patch

    from noodle.artifacts import write_bytes
    from noodle_nodes.document_intelligence import pdf_extract_text

    fake_ref = write_bytes(b"fake", name="test.pdf", content_type="application/pdf")

    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Hello world this is page one."
    mock_page.page_number = 1

    mock_pdf_ctx = MagicMock()
    mock_pdf_ctx.__enter__ = lambda s: s
    mock_pdf_ctx.__exit__ = MagicMock(return_value=False)
    mock_pdf_ctx.pages = [mock_page]

    with patch("pdfplumber.open", return_value=mock_pdf_ctx):
        result = pdf_extract_text(input=fake_ref)

    assert "Hello world" in result["text"]
    assert result["pages"] == 1
    assert result["pages_extracted"] == 1
    assert is_artifact_ref(result["artifact"])


def test_pdf_extract_tables_raises_without_input(store_ctx) -> None:
    from noodle_nodes.document_intelligence import pdf_extract_tables
    with pytest.raises(ValueError, match="input is required"):
        pdf_extract_tables(input=None)


def test_pdf_extract_tables_no_tables_returns_zero(store_ctx) -> None:
    pytest.importorskip("pdfplumber")
    pytest.importorskip("pandas")
    from unittest.mock import MagicMock, patch

    from noodle.artifacts import write_bytes
    from noodle_nodes.document_intelligence import pdf_extract_tables

    fake_ref = write_bytes(b"fake", name="test.pdf", content_type="application/pdf")

    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Some text but no tables"
    mock_page.extract_tables.return_value = []

    mock_pdf_ctx = MagicMock()
    mock_pdf_ctx.__enter__ = lambda s: s
    mock_pdf_ctx.__exit__ = MagicMock(return_value=False)
    mock_pdf_ctx.pages = [mock_page]

    with patch("pdfplumber.open", return_value=mock_pdf_ctx):
        result = pdf_extract_tables(input=fake_ref)

    assert result["tables_found"] == 0
    assert result["dataset"] is None


def test_pdf_extract_tables_returns_dataset_ref(store_ctx) -> None:
    pytest.importorskip("pdfplumber")
    pytest.importorskip("pandas")
    from unittest.mock import MagicMock, patch

    from noodle.artifacts import write_bytes
    from noodle_nodes.document_intelligence import pdf_extract_tables

    fake_ref = write_bytes(b"fake", name="test.pdf", content_type="application/pdf")

    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Some text"
    mock_page.extract_tables.return_value = [
        [["Name", "Age"], ["Alice", "30"], ["Bob", "25"]]
    ]

    mock_pdf_ctx = MagicMock()
    mock_pdf_ctx.__enter__ = lambda s: s
    mock_pdf_ctx.__exit__ = MagicMock(return_value=False)
    mock_pdf_ctx.pages = [mock_page]

    with patch("pdfplumber.open", return_value=mock_pdf_ctx):
        result = pdf_extract_tables(input=fake_ref)

    assert result["tables_found"] == 1
    assert is_dataset_ref(result["dataset"])
    assert result["summary"][0]["rows"] == 2
    assert result["summary"][0]["columns"] == 2


def test_parse_page_selection_blank_returns_all(store_ctx) -> None:
    from noodle_nodes.document_intelligence import _parse_page_selection
    assert _parse_page_selection("", 5, 0) == [0, 1, 2, 3, 4]


def test_parse_page_selection_range(store_ctx) -> None:
    from noodle_nodes.document_intelligence import _parse_page_selection
    assert _parse_page_selection("1-3", 5, 0) == [0, 1, 2]


def test_parse_page_selection_list(store_ctx) -> None:
    from noodle_nodes.document_intelligence import _parse_page_selection
    assert _parse_page_selection("1,3,5", 5, 0) == [0, 2, 4]


def test_pdf_generate_raises_without_template(store_ctx) -> None:
    from noodle_nodes.document_intelligence import pdf_generate
    with pytest.raises(ValueError, match="template is required"):
        pdf_generate(input={}, template="")


def test_pdf_generate_raises_missing_package(store_ctx, monkeypatch) -> None:
    import builtins
    real_import = builtins.__import__

    def mock_import(name, *args, **kwargs):
        if name in ("weasyprint", "jinja2"):
            raise ImportError(f"No module named '{name}'")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", mock_import)
    from noodle_nodes.document_intelligence import pdf_generate
    with pytest.raises(RuntimeError, match="weasyprint"):
        pdf_generate(input={}, template="<h1>hi</h1>")


def test_pdf_generate_returns_artifact(store_ctx) -> None:
    pytest.importorskip("weasyprint")
    pytest.importorskip("jinja2")
    from noodle_nodes.document_intelligence import pdf_generate

    result = pdf_generate(
        input={"title": "Test Report", "body": "Hello world"},
        template="<html><body><h1>{{ title }}</h1><p>{{ body }}</p></body></html>",
        filename="test.pdf",
    )
    assert is_artifact_ref(result["artifact"])
    assert result["artifact"]["content_type"] == "application/pdf"
    assert result["size_bytes"] > 0


def test_pdf_generate_renders_template_variables(store_ctx) -> None:
    pytest.importorskip("weasyprint")
    pytest.importorskip("jinja2")
    from noodle.artifacts import read_bytes
    from noodle_nodes.document_intelligence import pdf_generate

    result = pdf_generate(
        input={"name": "Alice"},
        template="<html><body><p>Hello {{ name }}</p></body></html>",
    )
    # PDF is binary — just verify it's a non-empty valid artifact
    assert is_artifact_ref(result["artifact"])
    raw = read_bytes(result["artifact"])
    assert raw[:4] == b"%PDF"


def test_pdf_generate_blocks_external_resource_fetches(store_ctx, monkeypatch) -> None:
    from noodle_nodes.document_intelligence import pdf_generate

    class FakeHTML:
        def __init__(self, *, string: str, url_fetcher) -> None:
            self.string = string
            self.url_fetcher = url_fetcher

        def write_pdf(self) -> bytes:
            return self.url_fetcher("file:///etc/passwd")  # type: ignore[return-value]

    monkeypatch.setitem(sys.modules, "weasyprint", SimpleNamespace(HTML=FakeHTML))

    with pytest.raises(ValueError, match="external resource fetch blocked"):
        pdf_generate(
            input={},
            template='<html><body><img src="file:///etc/passwd"></body></html>',
        )


def test_docx_generate_raises_missing_package(store_ctx, monkeypatch) -> None:
    import builtins
    real_import = builtins.__import__

    def mock_import(name, *args, **kwargs):
        if name == "docx":
            raise ImportError("No module named 'docx'")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", mock_import)
    from noodle_nodes.document_intelligence import docx_generate
    with pytest.raises(RuntimeError, match="python-docx"):
        docx_generate(input={"name": "Alice"})


def test_docx_generate_returns_artifact(store_ctx) -> None:
    pytest.importorskip("docx")
    from noodle_nodes.document_intelligence import docx_generate

    result = docx_generate(input={"greeting": "Hello"}, filename="test.docx")
    assert is_artifact_ref(result["artifact"])
    assert "wordprocessingml" in result["artifact"]["content_type"]


def test_docx_generate_fills_template_placeholders(store_ctx) -> None:
    pytest.importorskip("docx")
    from docx import Document

    from noodle.artifacts import read_bytes, write_bytes
    from noodle_nodes.document_intelligence import docx_generate

    # Build a real .docx template with {{name}} placeholder
    doc = Document()
    doc.add_paragraph("Dear {{name}},")
    doc.add_paragraph("Your order {{order_id}} is ready.")
    buf = io.BytesIO()
    doc.save(buf)
    template_ref = write_bytes(
        buf.getvalue(),
        name="template.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    result = docx_generate(
        input={"name": "Alice", "order_id": "ORD-001"},
        template_artifact=template_ref,
        filename="filled.docx",
    )
    assert is_artifact_ref(result["artifact"])
    assert result["placeholders_filled"] == 2

    # Verify the text was actually replaced
    filled_bytes = read_bytes(result["artifact"])
    filled_doc = Document(io.BytesIO(filled_bytes))
    full_text = " ".join(p.text for p in filled_doc.paragraphs)
    assert "Alice" in full_text
    assert "ORD-001" in full_text
    assert "{{name}}" not in full_text


def test_docx_generate_raises_for_unfilled_placeholders(store_ctx) -> None:
    pytest.importorskip("docx")
    from docx import Document

    from noodle.artifacts import write_bytes
    from noodle_nodes.document_intelligence import docx_generate

    doc = Document()
    doc.add_paragraph("Hello {{name}}, your code is {{code}}.")
    buf = io.BytesIO()
    doc.save(buf)
    template_ref = write_bytes(
        buf.getvalue(),
        name="tpl.docx",
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )

    with pytest.raises(ValueError, match="name"):
        docx_generate(input={}, template_artifact=template_ref)


def test_docx_extract_returns_paragraphs(store_ctx) -> None:
    pytest.importorskip("docx")
    from docx import Document

    from noodle.artifacts import write_bytes
    from noodle_nodes.document_intelligence import docx_extract

    doc = Document()
    doc.add_heading("My Title", level=1)
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    buf = io.BytesIO()
    doc.save(buf)
    artifact_ref = write_bytes(
        buf.getvalue(),
        name="test.docx",
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )

    result = docx_extract(input=artifact_ref)
    assert result["paragraph_count"] >= 3
    headings = [p for p in result["paragraphs"] if p["is_heading"]]
    assert any("My Title" in h["text"] for h in headings)


def test_docx_extract_returns_tables(store_ctx) -> None:
    pytest.importorskip("docx")
    from docx import Document

    from noodle.artifacts import write_bytes
    from noodle_nodes.document_intelligence import docx_extract

    doc = Document()
    tbl = doc.add_table(rows=3, cols=2)
    tbl.cell(0, 0).text = "Name"
    tbl.cell(0, 1).text = "Score"
    tbl.cell(1, 0).text = "Alice"
    tbl.cell(1, 1).text = "95"
    tbl.cell(2, 0).text = "Bob"
    tbl.cell(2, 1).text = "87"
    buf = io.BytesIO()
    doc.save(buf)
    artifact_ref = write_bytes(
        buf.getvalue(),
        name="test.docx",
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )

    result = docx_extract(input=artifact_ref, include_tables=True)
    assert result["table_count"] == 1
    assert result["tables"][0]["headers"] == ["Name", "Score"]
    assert result["tables"][0]["row_count"] == 2


def test_excel_report_generate_raises_without_input(store_ctx) -> None:
    from noodle_nodes.document_intelligence import excel_report_generate
    with pytest.raises(ValueError, match="input is required"):
        excel_report_generate(input=None)


def test_excel_report_generate_raises_on_empty_dataset(store_ctx) -> None:
    pytest.importorskip("openpyxl")
    pytest.importorskip("pandas")
    from noodle_nodes.document_intelligence import excel_report_generate
    with pytest.raises(ValueError, match="empty"):
        excel_report_generate(input=[])


def test_excel_report_generate_returns_artifact_from_records(store_ctx) -> None:
    pytest.importorskip("openpyxl")
    pytest.importorskip("pandas")
    from noodle.artifacts import read_bytes
    from noodle_nodes.document_intelligence import excel_report_generate

    result = excel_report_generate(
        input=[{"name": "Alice", "score": 95}, {"name": "Bob", "score": 87}],
        sheet_name="Results",
        title="Test Results",
        filename="results.xlsx",
    )
    assert is_artifact_ref(result["artifact"])
    assert "spreadsheet" in result["artifact"]["content_type"]
    assert result["rows"] == 2
    assert result["columns"] == 2

    raw = read_bytes(result["artifact"])
    assert raw[:4] == b"PK\x03\x04"  # xlsx is a zip file


def test_excel_report_generate_returns_artifact_from_dataset(store_ctx) -> None:
    pytest.importorskip("openpyxl")
    pytest.importorskip("pandas")
    from noodle_nodes.datasets import records_to_dataset
    from noodle_nodes.document_intelligence import excel_report_generate

    ds = records_to_dataset([{"x": 1, "y": 2}, {"x": 3, "y": 4}])
    result = excel_report_generate(input=ds, filename="out.xlsx")
    assert is_artifact_ref(result["artifact"])


def test_excel_extract_returns_dataset_ref(store_ctx) -> None:
    pytest.importorskip("openpyxl")
    pytest.importorskip("pandas")
    from noodle_nodes.document_intelligence import excel_extract, excel_report_generate

    # Generate an xlsx first, then extract it
    gen_result = excel_report_generate(
        input=[{"fruit": "apple", "count": 5}, {"fruit": "banana", "count": 3}],
        filename="fruit.xlsx",
    )
    result = excel_extract(input=gen_result["artifact"])
    assert is_dataset_ref(result)


def test_excel_extract_raises_without_input(store_ctx) -> None:
    from noodle_nodes.document_intelligence import excel_extract
    with pytest.raises(ValueError, match="input is required"):
        excel_extract(input=None)


def test_barcode_qr_generate_raises_without_input(store_ctx) -> None:
    from noodle_nodes.document_intelligence import barcode_qr_generate
    with pytest.raises(ValueError, match="input is required"):
        barcode_qr_generate(input=None)


def test_barcode_qr_generate_qr_returns_png_artifact(store_ctx) -> None:
    pytest.importorskip("qrcode")
    from noodle_nodes.document_intelligence import barcode_qr_generate

    result = barcode_qr_generate(input="https://example.com", format="qr", filename="qr.png")
    assert is_artifact_ref(result["artifact"])
    assert result["artifact"]["content_type"] == "image/png"
    assert result["format"] == "qr"
    assert result["value"] == "https://example.com"


def test_barcode_qr_decode_raises_without_input(store_ctx) -> None:
    from noodle_nodes.document_intelligence import barcode_qr_decode
    with pytest.raises(ValueError, match="input is required"):
        barcode_qr_decode(input=None)


def test_barcode_qr_decode_requirements_are_platform_adaptive() -> None:
    from noodle.sdk import registry
    manifest = next(m for m in registry.manifests() if m.id == "barcode_qr_decode")
    reqs = manifest.requirements
    assert any("zxing" in r and "win32" in r for r in reqs), "zxing-cpp win32 req missing"
    assert any("pyzbar" in r and "win32" in r for r in reqs), "pyzbar non-win32 req missing"


def test_barcode_qr_roundtrip(store_ctx) -> None:
    """Generate a QR code then decode it — values must match."""
    import sys
    pytest.importorskip("qrcode")
    if sys.platform == "win32":
        pytest.importorskip("zxing_cpp")
    else:
        pytest.importorskip("pyzbar")
    from noodle_nodes.document_intelligence import barcode_qr_decode, barcode_qr_generate

    gen_result = barcode_qr_generate(input="HELLO-NOODLE-123", format="qr")
    decode_result = barcode_qr_decode(input=gen_result["artifact"])

    assert decode_result["count"] == 1
    assert decode_result["codes"][0]["data"] == "HELLO-NOODLE-123"


def test_barcode_qr_decode_raises_on_blank_image(store_ctx) -> None:
    import sys
    if sys.platform == "win32":
        pytest.importorskip("zxing_cpp")
    else:
        pytest.importorskip("pyzbar")
    pytest.importorskip("PIL")
    from PIL import Image

    from noodle.artifacts import write_bytes
    from noodle_nodes.document_intelligence import barcode_qr_decode

    # Blank white image — no barcode
    img = Image.new("RGB", (100, 100), color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    ref = write_bytes(buf.getvalue(), name="blank.png", content_type="image/png")

    with pytest.raises(ValueError, match="No barcode"):
        barcode_qr_decode(input=ref)


def test_document_intelligence_nodes_registered() -> None:
    from noodle.sdk import registry
    ids = {m.id for m in registry.manifests()}
    expected = {
        "pdf_extract_text",
        "pdf_extract_tables",
        "pdf_generate",
        "docx_generate",
        "docx_extract",
        "excel_report_generate",
        "excel_extract",
        "barcode_qr_generate",
        "barcode_qr_decode",
    }
    assert expected <= ids, f"Missing from registry: {expected - ids}"


def test_document_intelligence_nodes_have_requirements() -> None:
    """Every node with optional packages must declare requirements."""
    from noodle.sdk import registry
    nodes_with_reqs = {
        "pdf_extract_text",
        "pdf_extract_tables",
        "pdf_generate",
        "docx_generate",
        "docx_extract",
        "excel_report_generate",
        "excel_extract",
        "barcode_qr_generate",
        "barcode_qr_decode",
    }
    manifests = {m.id: m for m in registry.manifests()}
    for node_id in nodes_with_reqs:
        manifest = manifests.get(node_id)
        assert manifest is not None, f"Node {node_id} not in registry"
        assert manifest.requirements, f"Node {node_id} has no requirements declared"


def test_import_does_not_import_optional_packages() -> None:
    """Importing noodle_nodes must not pull in any optional document packages."""
    # Check in a fresh interpreter so earlier tests that intentionally exercise
    # optional Excel/PDF features cannot pollute this assertion via sys.modules.
    code = """
import json
import sys
import noodle_nodes  # noqa: F401

forbidden = {
    "pdfplumber", "weasyprint", "docx", "openpyxl",
    "qrcode", "barcode", "pyzbar",
}
print(json.dumps(sorted(forbidden & set(sys.modules))))
"""
    result = subprocess.run(
        [sys.executable, "-c", code],
        check=True,
        capture_output=True,
        text=True,
    )
    leaked = set(json.loads(result.stdout))
    assert not leaked, f"Optional packages leaked into module scope: {leaked}"


def test_document_intelligence_importable_without_optional_packages() -> None:
    """Module must import cleanly even when no doc-processing packages are installed."""
    import importlib
    mod = importlib.import_module("noodle_nodes.document_intelligence")
    assert hasattr(mod, "pdf_extract_text")
    assert hasattr(mod, "pdf_extract_tables")
    assert hasattr(mod, "pdf_generate")
    assert hasattr(mod, "docx_generate")
    assert hasattr(mod, "excel_report_generate")
    assert hasattr(mod, "docx_extract")
    assert hasattr(mod, "excel_extract")
    assert hasattr(mod, "barcode_qr_generate")
    assert hasattr(mod, "barcode_qr_decode")
